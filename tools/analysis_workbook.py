"""One spreadsheet from a batch of finished runs.

``simulations/<run>/analysis/`` holds the notebook's output as text and PNGs,
one folder per run. That is the right shape for reading a single run and the
wrong shape for comparing twenty: the numbers you want side by side are spread
over four files in twenty folders, and nobody is going to retype them.

This module parses those text files back into numbers and writes one
``.xlsx`` with a row per run — the same layout as the hand-made
``crosslink-analysis-summary.xlsx``:

``Summary``
    one row per run: what the run was, whether its data is usable, the K-K
    *contact* statistics (pairs that came within the cutoff) and the K-K *bond*
    statistics (crosslinks that actually formed), plus surface bonds.
``Contacts detail``
    the full contact breakdown, including the per-pair probabilities and the
    closest approach that the Summary only samples.
``Equilibration``
    burn-in, correlation time and N_eff per observable — the numbers that say
    whether any of the averages in the other two sheets mean anything.
``Notes``
    what the columns are, and the caveats that travel with them.

Contacts and bonds answer different questions and are easy to confuse, which is
why they are separate column groups rather than interleaved. A *contact* is two
lysine beads within the analysis cutoff in a saved frame; it is counted every
frame and nothing is consumed. A *bond* is a crosslink the reactive run
actually formed; it happens once, it is permanent, and it uses up both lysines'
valence. A run with many contacts and few bonds is telling you the reaction was
rate-limited, not that the chains never met.

Parsing text that was written for humans is fragile, so every field is pulled
by a named regex and a field that does not match comes back as None and is left
blank rather than guessed at. A run whose analysis was never generated still
gets a row, marked in the "Data quality" column, because a missing run is
information too.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SIMULATIONS = Path("simulations")

# Runs whose trajectories are known to be physically meaningless, and why. They
# still get a row — leaving them out silently would make the workbook look like
# the batch succeeded — but the "Data quality" column says not to use them.
KNOWN_BAD: dict[str, str] = {}


def _num(text: str | None) -> float | int | None:
    """'1,234' or '12.3' -> a number; None for anything that isn't one."""
    if text is None:
        return None
    cleaned = text.replace(",", "").strip()
    if not cleaned:
        return None
    try:
        value = float(cleaned)
    except ValueError:
        return None
    return int(value) if value.is_integer() and "." not in cleaned else value


def _find(pattern: str, text: str, group: int = 1) -> str | None:
    match = re.search(pattern, text)
    return match.group(group) if match else None


def _ratio(a: Any, b: Any) -> float | None:
    """a / b, rounded, or None when either is missing or b is zero."""
    if a is None or b is None or not b:
        return None
    return round(a / b, 4)


def _share(part: Any, other: Any) -> float | None:
    """part / (part + other), the fraction of the total — None if undefined."""
    if part is None or other is None or not (part + other):
        return None
    return round(part / (part + other), 4)


@dataclass
class RunRow:
    """Everything the workbook knows about one run."""

    name: str
    family: str = ""
    mode: str = ""
    crosslinking: str = ""
    quality: str = "OK"
    values: dict[str, Any] = field(default_factory=dict)

    def get(self, key: str) -> Any:
        return self.values.get(key)


def split_name(name: str) -> tuple[str, str, str]:
    """``triblock-64-pre-xl-v3`` -> ('triblock-64', 'preattached', 'on').

    The family is whatever precedes the mode, so it keeps working for names
    like ``triblock-32`` that contain a hyphen of their own. Trailing tags are
    stripped first: a version (``-v3``) and a concentration marker
    (``-c010`` = 0.010 chains/nm^2) and a preattached-fraction marker
    (``-f30`` = 0.3 of all lysines pinned) and a run-length marker
    (``-t600`` = 600 ns) and a bare batch label such as ``-massmatch``,
    in any order and any number,
    so adding a tag to a run name does not cost it its family.
    """
    stem = name
    while True:
        stripped = re.sub(r"-(?:v\d+|c\d+|f\d+|t\d+|mm\d+|massmatch)$", "", stem)
        if stripped == stem:
            break
        stem = stripped
    match = re.search(r"-(brush|free|pre)-(xl|noxl)$", stem)
    if not match:
        return stem, "", ""
    family = stem[: match.start()]
    mode = {"brush": "brush", "free": "free", "pre": "preattached"}[match.group(1)]
    return family, mode, "on" if match.group(2) == "xl" else "off"


def parse_contacts(text: str) -> dict[str, Any]:
    """The numbers in ``crosslink_contacts_K.txt``."""
    out: dict[str, Any] = {}
    out["cutoff_A"] = _num(_find(r"within ([\d.]+) Å", text))
    out["lysines"] = _num(_find(r"beads tracked:\s+(\d+)", text))
    out["lys_per_chain"] = _num(_find(r"\((\d+) per chain", text))
    out["chains"] = _num(_find(r"per chain x (\d+) chains", text))
    out["frames_analysed"] = _num(_find(r"frames analysed:\s+(\d+)/", text))
    out["frames_total"] = _num(_find(r"frames analysed:\s+\d+/(\d+)", text))
    out["from_frame"] = _num(_find(r"from frame (\d+)", text))
    out["pairs_intra"] = _num(_find(r"pairs monitored:\s+(\d+) intra", text))
    out["pairs_inter"] = _num(_find(r"\+ (\d+) inter-chain", text))
    out["hits_intra"] = _num(_find(r"hits \(pair-frames\)\s+([\d,]+)", text))
    out["hits_inter"] = _num(_find(r"hits \(pair-frames\)\s+[\d,]+\s+([\d,]+)", text))
    out["hits_per_frame_intra"] = _num(_find(r"hits per frame\s+([\d.]+)", text))
    out["hits_per_frame_inter"] = _num(_find(r"hits per frame\s+[\d.]+\s+([\d.]+)", text))
    out["p_intra"] = _num(_find(r"P\(pair in contact\)\s+([\d.]+)", text))
    out["p_inter"] = _num(_find(r"P\(pair in contact\)\s+[\d.]+\s+([\d.]+)", text))
    out["ever_intra"] = _num(_find(r"pairs ever in contact\s+(\d+)/", text))
    out["ever_inter"] = _num(_find(r"pairs ever in contact\s+\d+/\d+\s+(\d+)/", text))
    out["pairs_pinned_dropped"] = _num(
        _find(r"\((\d+) pairs dropped: both lysines surface-bonded", text)) or 0
    out["pairs_seqsep_dropped"] = _num(
        _find(r"\((\d+) intra-chain pairs dropped: <", text)) or 0
    out["closest_intra_A"] = _num(_find(r"closest approach \(Å\)\s+([\d.]+)", text))
    out["closest_inter_A"] = _num(_find(r"closest approach \(Å\)\s+[\d.]+\s+([\d.]+)", text))
    return out


def parse_bonds(text: str) -> dict[str, Any]:
    """The numbers in ``crosslinks_formed.txt``."""
    out: dict[str, Any] = {}
    out["bond_mode"] = _find(r"\(mode (\w+),", text)
    out["frames"] = _num(_find(r"\(mode \w+, (\d+) frames", text))
    out["saved_ns"] = _num(_find(r"frames, ([\d.]+) ns\)", text))
    out["lysines"] = _num(_find(r"lysines in the system:\s+(\d+)", text))
    out["kk_total"] = _num(_find(r"K–K crosslinks:\s+(\d+)", text))
    out["kk_lys_used"] = _num(_find(r"K–K crosslinks:\s+\d+\s+using (\d+) lysines", text))
    out["kk_lys_used_pct"] = _num(_find(r"K–K crosslinks:.*?\(([\d.]+)%\)", text))
    out["kk_intra"] = _num(_find(r"intra-chain \(loops\):\s+(\d+)", text))
    out["kk_inter"] = _num(_find(r"inter-chain \(bridges\):\s+(\d+)", text))
    out["surf_total"] = _num(_find(r"K–surface bonds \(ground\):\s+(\d+)", text))
    out["surf_initial"] = _num(_find(r"present at t=0:\s+(\d+)", text))
    out["surf_run"] = _num(_find(r"formed during the run:\s+(\d+)", text))
    out["unreacted"] = _num(_find(r"lysines still unreacted:\s+(\d+)", text))
    out["first_bond_ns"] = _num(_find(r"first bond formed at\s+([\d.]+) ns", text))
    out["last_bond_ns"] = _num(_find(r"last at ([\d.]+) ns", text))
    out["half_bonds_ns"] = _num(_find(r"half the K–K bonds by\s+([\d.]+) ns", text))
    out["chains_bridged"] = _num(_find(r"chains bridged to another:\s+(\d+) of", text))
    out["chains_total"] = _num(_find(r"chains bridged to another:\s+\d+ of (\d+)", text))
    out["largest_group"] = _num(_find(r"largest connected group (\d+)", text))
    return out


def parse_equilibration(text: str) -> dict[str, Any]:
    """The numbers in ``equilibration.txt``."""
    out: dict[str, Any] = {}
    out["eq_frames"] = _num(_find(r"trajectory:\s+(\d+) frames", text))
    out["eq_chains"] = _num(_find(r"(\d+) chains x", text))
    out["eq_residues"] = _num(_find(r"chains x (\d+) residues", text))
    out["eq_total_ns"] = _num(_find(r"residues, ([\d.]+) ns total", text))
    for label, key in (("Rg", "rg"), ("RMSD", "rmsd")):
        row = rf"{label}\s+([\d.]+) ns\s+([\d.]+) ns\s+(\d+)\s+([\d.]+) \+/- ([\d.]+)\s+([\d.]+)"
        match = re.search(row, text)
        if match:
            out[f"{key}_eq_ns"] = _num(match.group(1))
            out[f"{key}_tau_ns"] = _num(match.group(2))
            out[f"{key}_neff"] = _num(match.group(3))
            out[f"{key}_plateau_nm"] = _num(match.group(4))
            out[f"{key}_err_nm"] = _num(match.group(5))
            out[f"{key}_drift_sigma"] = _num(match.group(6))
    out["burnin_frames"] = _num(_find(r"burn-in to discard: (\d+) frames", text))
    out["burnin_ns"] = _num(_find(r"burn-in to discard: \d+ frames \(([\d.]+) ns\)", text))
    out["burnin_pct"] = _num(_find(r"= (\d+)% of the run", text))
    out["next_steps"] = _num(_find(r"next run of this sequence: ~([\d,]+) steps", text))
    return out


def parse_metadata(runtime: Path) -> dict[str, Any]:
    """The few fields the workbook wants out of ``runtime/metadata.csv``."""
    path = runtime / "metadata.csv"
    if not path.is_file():
        return {}
    wanted = {"steps_per_frame": "steps_per_frame", "total_steps": "total_steps",
              "n_frames_saved": "meta_frames", "nmol": "meta_chains",
              "n_residues": "meta_residues", "status": "status",
              "spacing_nm": "spacing_nm", "concentration_chains_per_nm2": "concentration",
              "crosslink_distance_nm": "xl_distance_nm", "crosslink_prob": "xl_prob",
              "surface_preattached_fraction": "preattached_fraction",
              "surface_max_fraction": "surface_max_fraction",
              "run_wall_seconds": "wall_seconds"}
    out: dict[str, Any] = {}
    with open(path, newline="") as f:
        for row in csv.DictReader(f):
            key = wanted.get(row.get("key", ""))
            if key:
                out[key] = _num(row.get("value")) if key != "status" else row.get("value")
    return out


def collect(name: str) -> RunRow:
    """Read one run's analysis folder into a row, however complete it is."""
    family, mode, crosslinking = split_name(name)
    row = RunRow(name=name, family=family, mode=mode, crosslinking=crosslinking)
    analysis = SIMULATIONS / name / "analysis"
    runtime = SIMULATIONS / name / "runtime"

    row.values.update(parse_metadata(runtime))
    found = []
    for filename, parser in (("crosslink_contacts_K.txt", parse_contacts),
                             ("crosslinks_formed.txt", parse_bonds),
                             ("equilibration.txt", parse_equilibration)):
        path = analysis / filename
        if path.is_file():
            row.values.update(parser(path.read_text()))
            found.append(filename.split(".")[0])

    if name in KNOWN_BAD:
        row.quality = KNOWN_BAD[name]
    elif not found:
        row.quality = "no analysis — run the notebook's Analyze section"
    elif "crosslink_contacts_K" not in found:
        row.quality = "OK (no contacts analysis)"
    return row


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------

SUMMARY_GROUPS = [
    ("the run", 10),
    ("K–K contacts (pairs within the cutoff, counted every frame)", 7),
    ("K–K bonds actually formed (crosslinks)", 5),
    ("K–surface bonds", 3),
    ("equilibration", 4),
]

SUMMARY_COLUMNS = [
    ("Run", lambda r: r.name),
    ("Family", lambda r: r.family),
    ("Mode", lambda r: r.mode),
    ("Crosslinking", lambda r: r.crosslinking),
    ("Data quality", lambda r: r.quality),
    ("Frames", lambda r: r.get("frames") or r.get("eq_frames") or r.get("meta_frames")),
    ("Steps/frame", lambda r: r.get("steps_per_frame")),
    ("Saved time (ns)", lambda r: r.get("saved_ns") or r.get("eq_total_ns")),
    ("Chains", lambda r: r.get("chains") or r.get("eq_chains") or r.get("meta_chains")),
    ("Lysines", lambda r: r.get("lysines")),
    # contacts
    ("intra (pair-frames)", lambda r: r.get("hits_intra")),
    ("inter (pair-frames)", lambda r: r.get("hits_inter")),
    ("inter / intra", lambda r: _ratio(r.get("hits_inter"), r.get("hits_intra"))),
    ("inter share", lambda r: _share(r.get("hits_inter"), r.get("hits_intra"))),
    ("P(pair) intra", lambda r: r.get("p_intra")),
    ("P(pair) inter", lambda r: r.get("p_inter")),
    ("P inter / P intra", lambda r: _ratio(r.get("p_inter"), r.get("p_intra"))),
    # bonds
    ("K–K total", lambda r: r.get("kk_total")),
    ("intra (loops)", lambda r: r.get("kk_intra")),
    ("inter (bridges)", lambda r: r.get("kk_inter")),
    ("inter / intra", lambda r: _ratio(r.get("kk_inter"), r.get("kk_intra"))),
    ("inter share", lambda r: _share(r.get("kk_inter"), r.get("kk_intra"))),
    # surface
    ("total", lambda r: r.get("surf_total")),
    ("at t=0", lambda r: r.get("surf_initial")),
    ("formed in run", lambda r: r.get("surf_run")),
    # equilibration
    ("burn-in (ns)", lambda r: r.get("burnin_ns")),
    ("burn-in (% of run)", lambda r: r.get("burnin_pct")),
    ("N_eff (worst)", lambda r: _worst_neff(r)),
    ("usable? (N_eff>=50)", lambda r: _usable(r)),
]


def _worst_neff(row: RunRow) -> Any:
    values = [row.get("rg_neff"), row.get("rmsd_neff")]
    values = [v for v in values if v is not None]
    return min(values) if values else None


def _usable(row: RunRow) -> str | None:
    """The notebook's own threshold for an average worth quoting."""
    worst = _worst_neff(row)
    if worst is None:
        return None
    return "yes" if worst >= 50 else "no"


CONTACT_COLUMNS = [
    ("Run", lambda r: r.name),
    ("Frames analysed", lambda r: r.get("frames_analysed")),
    ("of frames", lambda r: r.get("frames_total")),
    ("From frame", lambda r: r.get("from_frame")),
    ("Cutoff (Å)", lambda r: r.get("cutoff_A")),
    ("Lysines", lambda r: r.get("lysines")),
    ("per chain", lambda r: r.get("lys_per_chain")),
    ("Chains", lambda r: r.get("chains")),
    ("pairs monitored intra", lambda r: r.get("pairs_intra")),
    ("pairs monitored inter", lambda r: r.get("pairs_inter")),
    ("hits intra", lambda r: r.get("hits_intra")),
    ("hits inter", lambda r: r.get("hits_inter")),
    ("inter / intra hits", lambda r: _ratio(r.get("hits_inter"), r.get("hits_intra"))),
    ("hits/frame intra", lambda r: r.get("hits_per_frame_intra")),
    ("hits/frame inter", lambda r: r.get("hits_per_frame_inter")),
    ("P(pair) intra", lambda r: r.get("p_intra")),
    ("P(pair) inter", lambda r: r.get("p_inter")),
    ("P inter / P intra", lambda r: _ratio(r.get("p_inter"), r.get("p_intra"))),
    ("pairs ever in contact intra", lambda r: r.get("ever_intra")),
    ("pairs ever in contact inter", lambda r: r.get("ever_inter")),
    ("ever-contacted intra", lambda r: _ratio(r.get("ever_intra"), r.get("pairs_intra"))),
    ("ever-contacted inter", lambda r: _ratio(r.get("ever_inter"), r.get("pairs_inter"))),
    ("closest intra (Å)", lambda r: r.get("closest_intra_A")),
    ("closest inter (Å)", lambda r: r.get("closest_inter_A")),
]

EQUILIBRATION_COLUMNS = [
    ("Run", lambda r: r.name),
    ("Family", lambda r: r.family),
    ("Mode", lambda r: r.mode),
    ("Crosslinking", lambda r: r.crosslinking),
    ("Frames", lambda r: r.get("eq_frames")),
    ("Residues/chain", lambda r: r.get("eq_residues")),
    ("Total (ns)", lambda r: r.get("eq_total_ns")),
    ("Rg equilibrated (ns)", lambda r: r.get("rg_eq_ns")),
    ("Rg tau (ns)", lambda r: r.get("rg_tau_ns")),
    ("Rg N_eff", lambda r: r.get("rg_neff")),
    ("Rg plateau (nm)", lambda r: r.get("rg_plateau_nm")),
    ("Rg error (nm)", lambda r: r.get("rg_err_nm")),
    ("RMSD equilibrated (ns)", lambda r: r.get("rmsd_eq_ns")),
    ("RMSD tau (ns)", lambda r: r.get("rmsd_tau_ns")),
    ("RMSD N_eff", lambda r: r.get("rmsd_neff")),
    ("RMSD plateau (nm)", lambda r: r.get("rmsd_plateau_nm")),
    ("burn-in (frames)", lambda r: r.get("burnin_frames")),
    ("burn-in (ns)", lambda r: r.get("burnin_ns")),
    ("burn-in (%)", lambda r: r.get("burnin_pct")),
    ("N_eff (worst)", _worst_neff),
    ("usable? (N_eff>=50)", _usable),
    ("steps suggested next", lambda r: r.get("next_steps")),
]

NOTES = [
    ("What is in this workbook", ""),
    ("", ""),
    ("Scope", "One row per run, read from simulations/<run>/analysis/. Generated by "
              "tools/analysis_workbook.py, so it can be regenerated after re-running "
              "the notebook's Analyze section."),
    ("", ""),
    ("Contacts vs. bonds — they answer different questions", ""),
    ("Contact", "Two lysine beads (LYS CA) within the analysis cutoff in a saved frame. "
                "Counted every frame, nothing is consumed, so one pair sitting together "
                "for 100 frames contributes 100 pair-frames."),
    ("Bond", "A crosslink the reactive run actually formed. Happens once, is permanent, "
             "and uses up valence at both lysines. Many contacts with few bonds means the "
             "reaction was rate-limited, not that the chains never met."),
    ("", ""),
    ("Reading the numbers", ""),
    ("inter share", "inter / (intra + inter). For contacts this is the share a bulk "
                    "experiment would see; for bonds it is the fraction of crosslinks that "
                    "bridge two different chains rather than looping one back on itself."),
    ("P(pair)", "Per-pair probability of being in contact in a given frame. There are far "
                "more inter-chain pairs than intra-chain ones, so a low P(pair) inter can "
                "still dominate the total."),
    ("intra (loops)", "Primary-loop estimate — the headline number for elastic network "
                      "theory. A loop carries no load."),
    ("N_eff", "Independent samples of the slowest observable. The notebook's threshold for "
              "an average worth quoting is 50; below that, treat plateau values as "
              "indicative and expect wide error bars."),
    ("", ""),
    ("Caveats that travel with every number here", ""),
    ("Valence", "At crosslink_valence = 1 every lysine is a degree-<=2 node, so the network "
                "has no branch points and its modulus is zero by construction. The topology "
                "is meaningful; a stiffness is not derivable yet. See docs/GAPS.md LIT-1."),
    ("Kinetics", "CALVADOS has no crosslinker, no solvent and no activation barrier, and its "
                 "Langevin friction is ~2700x below water. The sequence and topology of "
                 "events is meaningful; event times are not kinetics."),
    ("Sequences", "'G1 A1' in the run CSVs expands to VPGGGVPGAG..., not a literal GAGA "
                  "block. See docs/BUGS.md BUG-BATCH-2."),
]


def build_workbook(names: list[str], path: str | Path) -> Path:
    """Write the workbook for `names`. Returns the path written."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = [collect(name) for name in names]
    wb = Workbook()

    head_font = Font(bold=True)
    group_font = Font(bold=True, color="FFFFFF")
    group_fill = PatternFill("solid", fgColor="4C72B0")
    bad_fill = PatternFill("solid", fgColor="F2DCDB")

    def write_sheet(ws, columns, groups=None, freeze="B3" if True else None):
        start = 1
        if groups:
            col = 1
            for title, width in groups:
                cell = ws.cell(row=1, column=col, value=title)
                cell.font, cell.fill = group_font, group_fill
                if width > 1:
                    ws.merge_cells(start_row=1, start_column=col,
                                   end_row=1, end_column=col + width - 1)
                col += width
            start = 2
        for i, (title, _) in enumerate(columns, start=1):
            cell = ws.cell(row=start, column=i, value=title)
            cell.font = head_font
            cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        for r, row in enumerate(rows, start=start + 1):
            for i, (_, getter) in enumerate(columns, start=1):
                ws.cell(row=r, column=i, value=getter(row))
            if not row.quality.startswith("OK"):
                for i in range(1, len(columns) + 1):
                    ws.cell(row=r, column=i).fill = bad_fill
        for i, (title, _) in enumerate(columns, start=1):
            longest = max([len(str(title))] +
                          [len(str(getter(row) or "")) for _, getter in [columns[i - 1]]
                           for row in rows])
            ws.column_dimensions[get_column_letter(i)].width = min(38, max(9, longest + 2))
        ws.freeze_panes = ws.cell(row=start + 1, column=2)

    ws = wb.active
    ws.title = "Summary"
    write_sheet(ws, SUMMARY_COLUMNS, SUMMARY_GROUPS)
    write_sheet(wb.create_sheet("Contacts detail"), CONTACT_COLUMNS)
    write_sheet(wb.create_sheet("Equilibration"), EQUILIBRATION_COLUMNS)

    notes = wb.create_sheet("Notes")
    for r, (left, right) in enumerate(NOTES, start=1):
        a = notes.cell(row=r, column=1, value=left)
        if left and not right:
            a.font = head_font
        cell = notes.cell(row=r, column=2, value=right)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 110

    path = Path(path)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# A flat CSV of the contact statistics
# ---------------------------------------------------------------------------

CONTACT_CSV_COLUMNS = [
    ("run", lambda r: r.name),
    ("family", lambda r: r.family),
    ("residues_per_chain", lambda r: r.get("eq_residues")),
    ("concentration_chains_nm2", lambda r: r.get("concentration")),
    ("preattached_fraction", lambda r: r.get("preattached_fraction")),
    ("cutoff_A", lambda r: r.get("cutoff_A")),
    ("frames_analysed", lambda r: r.get("frames_analysed")),
    ("frames_total", lambda r: r.get("frames_total")),
    ("first_frame", lambda r: r.get("from_frame")),
    ("pairs_intra", lambda r: r.get("pairs_intra")),
    ("pairs_inter", lambda r: r.get("pairs_inter")),
    ("pairs_dropped_both_pinned", lambda r: r.get("pairs_pinned_dropped")),
    ("hits_intra", lambda r: r.get("hits_intra")),
    ("hits_inter", lambda r: r.get("hits_inter")),
    ("inter_share_of_hits", lambda r: _share(r.get("hits_inter"), r.get("hits_intra"))),
    ("hits_per_frame_intra", lambda r: r.get("hits_per_frame_intra")),
    ("hits_per_frame_inter", lambda r: r.get("hits_per_frame_inter")),
    ("p_pair_intra", lambda r: r.get("p_intra")),
    ("p_pair_inter", lambda r: r.get("p_inter")),
    ("p_intra_over_p_inter", lambda r: _ratio(r.get("p_intra"), r.get("p_inter"))),
    ("pairs_ever_in_contact_intra", lambda r: r.get("ever_intra")),
    ("pairs_ever_in_contact_inter", lambda r: r.get("ever_inter")),
    ("closest_intra_A", lambda r: r.get("closest_intra_A")),
    ("closest_inter_A", lambda r: r.get("closest_inter_A")),
]


def write_contacts_csv(names: list[str], path: str | Path, *,
                       delimiter: str = ";", decimal: str = ",") -> Path:
    """One row per run of contact statistics.

    The defaults are the European spreadsheet convention — ``;`` between fields
    and ``,`` inside numbers — because that is what the rest of this project's
    CSVs use and what a Dutch-locale Excel expects. It matters more than it
    looks: open a dot-decimal file in such an Excel and the dot is read as a
    *thousands* separator, so ``10.0`` becomes 100 and ``0.04`` becomes 4. Every
    number in the file is then silently wrong, with nothing to show for it.

    Pass ``delimiter=","`` and ``decimal="."`` for the US/pandas convention
    (``pd.read_csv`` with no arguments).
    """
    rows = [collect(name) for name in names]
    rows.sort(key=lambda r: (r.family, r.get("concentration") or 0,
                             r.get("preattached_fraction") or 0, r.name))

    def cell(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float):
            # repr, not a fixed format: these span 1e-05 to 1e5 and rounding to a
            # fixed number of decimals would flatten the small probabilities to 0.
            return repr(value).replace(".", decimal) if decimal != "." else repr(value)
        return str(value)

    path = Path(path)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=delimiter)
        writer.writerow([name for name, _ in CONTACT_CSV_COLUMNS])
        for row in rows:
            writer.writerow([cell(getter(row)) for _, getter in CONTACT_CSV_COLUMNS])
    return path
